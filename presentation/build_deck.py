"""Builds the seven-slide executive deck from the analysis outputs.
Run analysis/analysis.py first."""
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
CH = OUT / "charts"

NAVY = RGBColor(0x1B, 0x2A, 0x41)
ORANGE = RGBColor(0xE4, 0x57, 0x2E)
GREY = RGBColor(0x6B, 0x73, 0x80)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
FONT = "Calibri"

reg = pd.read_csv(OUT / "regional_summary.csv", index_col=0)
nat = pd.read_csv(OUT / "national_trend.csv", index_col=0, parse_dates=True).iloc[:, 0]
alloc = pd.read_csv(OUT / "allocation.csv", index_col=0)
cat = pd.read_csv(OUT / "category_priority_regions.csv", index_col=0)
chan = pd.read_csv(OUT / "channel_summary.csv", index_col=0)

n_now, n_prev = nat.iloc[-1], nat.iloc[-13]
below = reg[reg.gap_pts > 0]
prio = alloc.sort_values("allocation_sar_m", ascending=False)
najran = reg.loc["Najran"]
digital_cost = chan.loc[["Mobile App", "Web Portal"], "cost_per_txn_sar"].mean()
assisted_cost = chan.loc[["Branch", "Call Center"], "cost_per_txn_sar"].mean()

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, runs, size=16, color=NAVY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """runs: str or list of paragraphs; each paragraph str or list of (text, overrides) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        for chunk in (para if isinstance(para, list) else [(para, {})]):
            t, o = chunk if isinstance(chunk, tuple) else (chunk, {})
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def base(num, label, title, presenter, notes):
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, 0.18, 7.5, ORANGE)
    text(s, 0.55, 0.3, 9, 0.35, label.upper(), size=12, color=ORANGE, bold=True)
    text(s, 0.55, 0.62, 12.3, 1.1, title, size=28, bold=True)
    text(s, 0.55, 7.0, 9, 0.3, "Tayseer Digital Adoption  |  SDAIA Academy SDA-DSC-112 Capstone  |  Source: Tayseer services data, Jan 2022–Dec 2025",
         size=10, color=GREY)
    text(s, 11.3, 7.0, 1.6, 0.3, f"{num} / 7", size=10, color=GREY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = f"PRESENTER: {presenter}\n\n{notes}"
    return s


def kpi(slide, x, y, w, value, label, color=NAVY):
    box(slide, x, y, w, 1.35, LIGHT)
    text(slide, x + 0.15, y + 0.1, w - 0.3, 0.7, value, size=30, bold=True, color=color)
    text(slide, x + 0.15, y + 0.78, w - 0.3, 0.55, label, size=12, color=GREY)


regions_txt = ", ".join(prio.index[:-1]) + " and " + prio.index[-1]

# ---------- 1. BLUF / Ask ----------
s = base(1, "Bottom line  ·  the ask",
         f"Approve SAR 40M for four regions, {regions_txt}, to close the last gap to 65%",
         "Yousef Almutairi  (0:00–0:30)",
         "Open with the ask, inside 30 seconds:\n"
         f"\"We are asking you to approve SAR 40 million for four regions: {regions_txt}. "
         "Nationally Tayseer is already above 65%, but these four regions will not reach it on their own any time soon. "
         "Najran alone would need almost two years. Our recommendation puts every riyal where the gap won't close by itself.\"\n"
         "Then hand over to the situation.")
box(s, 0.55, 2.0, 12.25, 1.15, NAVY)
text(s, 0.8, 2.08, 11.8, 1.0,
     [[("Big idea:  ", {"color": ORANGE, "bold": True}),
       ("The national 65% target is met, but four regions are stuck below it. Put the full SAR 40M into those four, "
        "so all four reach it by Dec 2026 instead of as late as 2027.", {"color": WHITE})]],
     size=19, anchor=MSO_ANCHOR.MIDDLE)
xs = [0.55, 3.65, 6.75, 9.85]
for x, (r, row) in zip(xs, prio.iterrows()):
    box(s, x, 3.55, 2.95, 2.6, LIGHT)
    text(s, x + 0.2, 3.7, 2.6, 0.4, r, size=16, bold=True)
    text(s, x + 0.2, 4.15, 2.6, 0.8, f"SAR {row.allocation_sar_m:.0f}M", size=34, bold=True, color=ORANGE)
    text(s, x + 0.2, 5.0, 2.6, 1.1,
         [f"{row.adoption_pct:.1f}% today, gap {row.gap_pts:.1f} pts",
          f"{row.months_to_target:.0f} months to 65% at today's pace"], size=12, color=GREY)
text(s, 0.55, 6.35, 12.2, 0.5,
     [[("Decision requested today:  ", {"bold": True, "color": ORANGE}),
       ("release SAR 20M now and SAR 20M after a June 2026 checkpoint.", {})]], size=15)

# ---------- 2. Situation ----------
s = base(2, "Situation  ·  national status",
         f"Nationally, Tayseer is on track: {n_now:.1f}% digital adoption, above 65% since August 2025",
         "Yousef Almutairi  (0:30–1:15)",
         "Orient first: the line is national digital adoption by month, 2022 to 2025. The dashed orange line is the 65% target.\n"
         f"Then the message: adoption rose steadily from {nat.iloc[0]:.1f}% to {n_now:.1f}%, about +{n_now - n_prev:.1f} points a year, "
         "and crossed 65% in August 2025.\n"
         "Transition: \"If we stopped at the national number, we would say job done. But the average hides something.\"")
s.shapes.add_picture(str(CH / "01_national_trend.png"), Inches(0.55), Inches(1.95), width=Inches(9.3))
kpi(s, 10.2, 2.0, 2.65, f"{n_now:.1f}%", "National adoption, Dec 2025")
kpi(s, 10.2, 3.55, 2.65, f"+{n_now - n_prev:.1f} pts", "Change vs Dec 2024", GREEN)
kpi(s, 10.2, 5.1, 2.65, "Aug 2025", "Month the 65% target was crossed")

# ---------- 3. Complication ----------
s = base(3, "Complication  ·  regional gap",
         f"But {len(below)} of 13 regions are still below 65%, and the national average hides them",
         "Abdulwahab Alnassar  (1:15–2:00)",
         "Orient: each bar is a region's digital adoption in Dec 2025, sorted lowest to highest. Orange means below target, "
         "grey means on or above. The dashed line is 65%.\n"
         f"Message: {len(below)} of 13 regions are below target. Five large regions (Riyadh, Makkah, Eastern Province, Madinah, Qassim) "
         "pull the national average above 65%.\n"
         f"Najran is lowest at {najran.adoption_pct:.1f}%, more than 4 points short.\n"
         "Transition: \"But not every orange bar is equally urgent, and that matters for where the money goes.\"")
s.shapes.add_picture(str(CH / "02_regional_adoption.png"), Inches(0.55), Inches(1.85), height=Inches(5.05))
text(s, 9.9, 2.2, 3.0, 4.5,
     [[("8 regions", {"size": 30, "bold": True, "color": ORANGE})],
      "below the 65% target in Dec 2025",
      [("", {})],
      [(f"{najran.adoption_pct:.1f}%", {"size": 30, "bold": True, "color": ORANGE})],
      "Najran, the lowest region, is 4.1 pts short",
      [("", {})],
      [("5 regions", {"size": 30, "bold": True})],
      "carry the national average above target"], size=14, color=GREY)

# ---------- 4. Evidence ----------
c = cat.loc["Complaints"]
s = base(4, "Evidence  ·  which gaps matter",
         "Four regions will close the gap on their own within 3 months; four will not, and Najran needs 21",
         "Abdulwahab Alnassar  (2:00–4:00)",
         "Orient: this chart shows only the 8 below-target regions. Bar length is the number of months each region needs to reach 65% "
         "if it keeps its 2025 pace. The labels show the gap in points and the monthly pace.\n"
         "Message 1: Al-Jouf, Hail, Tabuk and Asir are within 0.7 points and will cross within about 3 months. They don't need money.\n"
         "Message 2: Jazan, Al-Baha, Northern Borders and Najran need 7 to 21 months. Najran has the biggest gap AND the slowest pace "
         "(+0.19 pts/month), so it gets worse treatment from time, not better.\n"
         f"Message 3 (box on the right): in these four regions, Complaints is the weakest service at {c.priority_regions:.1f}% digital, "
         f"against {c.national:.1f}% nationally. That tells us what the money should fix.\n"
         "Method: 2025 pace = linear trend of the last 12 months; adoption is user-weighted, the same as the Tableau calculated field.")
s.shapes.add_picture(str(CH / "03_months_to_target.png"), Inches(0.55), Inches(1.85), width=Inches(9.0))
box(s, 9.85, 2.0, 3.0, 4.6, LIGHT)
text(s, 10.05, 2.15, 2.65, 4.4,
     [[("Where the gap sits", {"bold": True, "size": 15})],
      [(f"{c.priority_regions:.1f}%", {"size": 30, "bold": True, "color": ORANGE})],
      f"Complaints digital adoption in the 4 priority regions (national: {c.national:.1f}%)",
      [("", {})],
      [("All 8", {"size": 22, "bold": True})],
      "service categories in these regions trail the national rate by 3–5 pts, so the problem is regional, not one product"],
     size=13, color=GREY)

# ---------- 5. Options ----------
s = base(5, "Options  ·  how to use SAR 40M",
         "We compared three ways to spend SAR 40M; only one aims every riyal at a gap that won't close by itself",
         "Firas Alnasser  (4:00–5:00)",
         "Walk through the three options left to right, using the same three questions for each.\n"
         "A, spread evenly across all 8 below-target regions at SAR 5M each: simple and seen as fair, but half the money goes to "
         "regions that cross 65% within about 3 months anyway, and Najran gets only SAR 5M.\n"
         "B, focus on the 4 structural laggers, weighted by the number of users that still need to switch: targets the real gap.\n"
         "C, one national channel-migration campaign: national is already at 66%, and most of the reach lands in large regions "
         "that are already above target.\n"
         "Transition: \"So we recommend option B.\"")
cols = ["Criteria", "A · Spread evenly", "B · Focus on 4 laggers", "C · National campaign"]
rows = [
    ("What it is", "SAR 5M to each of the 8 regions below 65%", "SAR 40M split across Najran, N. Borders, Al-Baha, Jazan by gap size",
     "One national push to move branch & call-centre users to app/web"),
    ("Hits the gap that won't self-close?", "Half: SAR 20M goes to regions closing in ≤3 months",
     "Yes: 100% to regions needing 7–21 months", "Mostly not: reach is driven by large, on-target regions"),
    ("Najran funding", "SAR 5M", f"SAR {alloc.loc['Najran', 'allocation_sar_m']:.0f}M", "Unclear / diluted"),
    ("Main risk", "Too thin to change pace anywhere", "Concentration: depends on local delivery", "Lifts the average again, leaves gaps"),
]
tbl = s.shapes.add_table(len(rows) + 1, 4, Inches(0.55), Inches(1.95), Inches(12.25), Inches(4.7)).table
widths = [2.6, 3.2, 3.25, 3.2]
for i, w in enumerate(widths):
    tbl.columns[i].width = Inches(w)
for r in range(len(rows) + 1):
    for ci in range(4):
        cell = tbl.cell(r, ci)
        val = cols[ci] if r == 0 else rows[r - 1][ci]
        cell.text = val
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.name = FONT
        para.runs[0].font.size = Pt(14 if r == 0 else 13)
        para.runs[0].font.bold = r == 0 or ci == 0
        highlight = ci == 2
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = ORANGE if highlight else NAVY
            para.runs[0].font.color.rgb = WHITE
        else:
            cell.fill.fore_color.rgb = RGBColor(0xFD, 0xEC, 0xE6) if highlight else (LIGHT if r % 2 else WHITE)
            para.runs[0].font.color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
for r in range(1, len(rows) + 1):
    tbl.rows[r].height = Inches(1.0)

# ---------- 6. Recommendation ----------
s = base(6, "Recommendation  ·  option B",
         "Recommend option B: weight the SAR 40M by gap and bring all four regions to 65% by December 2026",
         "Firas Alnasser  (5:00–6:15)",
         "Orient: left is the split; each region's share is proportional to the users who still need to switch to reach 65%.\n"
         f"What the money buys: (1) digitise Complaints and Permits, the weakest journeys; (2) assisted-digital desks in branches, "
         f"where adoption is lowest; (3) local awareness in these four regions.\n"
         f"Expected effect: Najran must roughly double its pace, from +{alloc.loc['Najran', 'monthly_pace_pts']:.2f} to "
         f"+{alloc.loc['Najran', 'pace_needed_by_dec26']:.2f} pts a month, to reach 65% by Dec 2026 instead of around Sep 2027. "
         "The other three reach it 3 to 6 months earlier than at today's pace.\n"
         f"Side benefit: digital transactions cost about SAR {digital_cost:.0f}, against SAR {assisted_cost:.0f} for branch or call centre, "
         "and score higher on CSAT.\n"
         "Uncertainty, say it plainly: the data shows where the gap is, not how many points SAR 1M buys. That is why we stage the release (next slide).")
s.shapes.add_picture(str(CH / "05_allocation.png"), Inches(0.55), Inches(1.95), width=Inches(6.4))
text(s, 7.3, 1.95, 5.6, 4.9,
     [[("What the money buys", {"bold": True, "size": 17})],
      [("1  ", {"bold": True, "color": ORANGE}), ("Digitise Complaints & Permits, the weakest journeys", {})],
      [("2  ", {"bold": True, "color": ORANGE}), (f"Assisted-digital desks in branches (branch adoption {chan.loc['Branch', 'adoption_pct']:.0f}% vs app {chan.loc['Mobile App', 'adoption_pct']:.0f}%)", {})],
      [("3  ", {"bold": True, "color": ORANGE}), ("Local awareness & onboarding in the four regions", {})],
      [("", {})],
      [("Expected effect", {"bold": True, "size": 17})],
      [("All 4 regions ≥ 65% by Dec 2026 ", {"bold": True}), ("(Najran: ~Sep 2027 at today's pace)", {"color": GREY})],
      [("Najran pace must rise ", {}), (f"+{alloc.loc['Najran', 'monthly_pace_pts']:.2f} → +{alloc.loc['Najran', 'pace_needed_by_dec26']:.2f} pts/month", {"bold": True})],
      [(f"Digital txn ≈ SAR {digital_cost:.0f} vs SAR {assisted_cost:.0f} assisted; higher CSAT", {"color": GREY})]],
     size=14)

# ---------- 7. Ask + next step ----------
half = (alloc.adoption_pct + alloc.gap_pts / 2)
s = base(7, "The ask  ·  next step",
         "Approve SAR 40M for the four regions today, released in two tranches with a June 2026 checkpoint",
         "Firas Alnasser  (6:15–7:00)",
         f"Restate the ask: \"We ask you to approve SAR 40 million for {regions_txt}: SAR 20 million now, "
         "and SAR 20 million in June 2026 if each region is at least halfway to 65%.\"\n"
         "The checkpoint values are on the slide and are tracked monthly on the Tableau dashboard.\n"
         "If a region misses its checkpoint, its second tranche is re-targeted, either to the region's weakest journey or to "
         "another region that is falling behind.\n"
         "Close: \"Nationally we have hit 65%. This decision makes sure every region does.\"")
box(s, 0.55, 1.95, 6.0, 4.75, NAVY)
text(s, 0.85, 2.1, 5.5, 4.5,
     [[("Decision today", {"color": ORANGE, "bold": True, "size": 16})],
      [("Approve SAR 40M", {"color": WHITE, "bold": True, "size": 32})],
      [(regions_txt, {"color": WHITE, "size": 16})],
      [("", {})],
      [("Tranche 1 · now", {"color": ORANGE, "bold": True})],
      [("SAR 20M to start the Complaints/Permits fixes and branch desks", {"color": WHITE})],
      [("Tranche 2 · June 2026", {"color": ORANGE, "bold": True})],
      [("SAR 20M, released region by region if the checkpoint is met", {"color": WHITE})]], size=15)
text(s, 6.95, 1.95, 5.9, 0.5, "June 2026 checkpoint (halfway to 65%)", size=17, bold=True)
tbl = s.shapes.add_table(5, 3, Inches(6.95), Inches(2.5), Inches(5.9), Inches(3.0)).table
hdr = ["Region", "Dec 2025", "Must reach by Jun 2026"]
for ci, h in enumerate(hdr):
    tbl.cell(0, ci).text = h
for ri, (r, row) in enumerate(alloc.sort_values("adoption_pct").iterrows(), start=1):
    for ci, v in enumerate([r, f"{row.adoption_pct:.1f}%", f"≥ {half[r]:.1f}%"]):
        tbl.cell(ri, ci).text = v
for ri in range(5):
    for ci in range(3):
        cell = tbl.cell(ri, ci)
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.name, run.font.size = FONT, Pt(14)
        run.font.bold = ri == 0 or ci == 2
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
        run.font.color.rgb = WHITE if ri == 0 else (ORANGE if ci == 2 else NAVY)
text(s, 6.95, 5.7, 5.9, 1.0,
     "Tracked monthly on the Tableau dashboard. A missed checkpoint sends that region's tranche 2 to its weakest journey or to the next lagging region.",
     size=13, color=GREY)

prs.save(ROOT / "presentation" / "Tayseer_Capstone_Executive_Story.pptx")
print("saved")
